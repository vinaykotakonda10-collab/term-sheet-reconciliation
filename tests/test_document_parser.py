from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_BREAK

from src.document_parser import DocumentParseError, UnsupportedDocumentError, compact_layout, parse_document
from src.evidence import EvidenceLocator

from .conftest import GENEL_PDF, IDBI_PDF


def test_pdf_parsing_keeps_pages_and_both_text_views(idbi_document, genel_document):
    assert idbi_document.file_type == "pdf" and idbi_document.page_count == 16
    assert genel_document.page_count == 13
    page3 = idbi_document.pages[2]
    assert "Coupon Rate  10.75% p.a." in page3.layout_text  # label and value on one line
    assert "Coupon Rate" in page3.reading_text and "10.75% p.a." in page3.reading_text
    assert len(idbi_document.sha256) == 64


def test_pdf_layout_view_aligns_table_rows(idbi_document):
    page8 = idbi_document.pages[7].layout_text
    assert "Issue Opening Date  September 29, 2014" in page8
    assert "Issue Closing Date  October 17, 2014" in page8


def test_compact_layout_trims_padding():
    assert compact_layout("   a      b  \n\n\n\n  c ") == "a  b\n\nc"


def test_unsupported_and_missing_files(tmp_path: Path):
    (tmp_path / "sheet.txt").write_text("x")
    with pytest.raises(UnsupportedDocumentError, match="unsupported file type"):
        parse_document(tmp_path / "sheet.txt")
    (tmp_path / "old.doc").write_bytes(b"x")
    with pytest.raises(UnsupportedDocumentError, match="legacy .doc"):
        parse_document(tmp_path / "old.doc")
    with pytest.raises(DocumentParseError, match="not found"):
        parse_document(tmp_path / "missing.pdf")


def test_corrupt_and_image_only_pdfs(tmp_path: Path):
    corrupt = tmp_path / "corrupt.pdf"
    corrupt.write_bytes(b"%PDF-1.4 this is not really a pdf")
    with pytest.raises(DocumentParseError):
        parse_document(corrupt)

    blank = tmp_path / "blank.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(blank)
    with pytest.raises(DocumentParseError, match="no extractable text"):
        parse_document(blank)


def _make_docx(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Term Sheet")
    doc.add_paragraph("ISIN: XS0000000009")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Coupon Rate", "5.50% p.a."
    table.cell(1, 0).text, table.cell(1, 1).text = "Maturity Date", "15 March 2030"
    doc.add_paragraph("Day count: 30/360").add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Minimum Subscription: EUR 100,000")
    doc.save(path)
    return path


def test_docx_parsing_paragraphs_tables_and_page_breaks(tmp_path: Path):
    parsed = parse_document(_make_docx(tmp_path / "sheet.docx"))
    assert parsed.file_type == "docx" and parsed.page_numbering == "page_breaks"
    assert parsed.page_count == 2
    assert "Coupon Rate | 5.50% p.a." in parsed.pages[0].layout_text
    assert "Minimum Subscription: EUR 100,000" in parsed.pages[1].reading_text
    locator = EvidenceLocator(parsed)
    assert locator.locate("Maturity Date | 15 March 2030").page == 1
    assert locator.locate("Minimum Subscription: EUR 100,000").page == 2


def test_corrupt_docx(tmp_path: Path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(DocumentParseError, match="not a readable Word document"):
        parse_document(bad)


@pytest.mark.parametrize("path", [IDBI_PDF, GENEL_PDF])
def test_sample_term_sheets_exist(path):
    assert path.exists()
